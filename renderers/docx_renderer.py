"""Word 渲染（python-docx）。

与 :mod:`renderers.pdf_renderer` 消费**同一套 Block**，
所以两种格式的题目顺序、题号、答案分组必然一致。

旧实现在 docx 层用 ``line[0].isdigit() and '. ' in line[:6]`` 从纯文本反解析题号 ——
这对**题干以数字开头的题**会误判。Block 中间表示从源头消除了这个隐患。

## 纸张

docx 没有"拼版"概念，A3 就按横向 A3 设页（内容整幅流排）。
「两张 A4 拼版」是 PDF 侧的特性（见 ``pdf_renderer``）。
"""

from __future__ import annotations

from typing import List, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from app.constants import PAPER_SIZES
from app.docmodel import (
    AnswerCardEssay,
    AnswerCardHead,
    AnswerCardSection,
    AnswerKey,
    Block,
    Heading,
    Instructions,
    PageBreak,
    Paragraph,
    QuestionBlock,
    SectionHeader,
    Spacer,
)

from .answer_card import bands

BODY_FONT = "宋体"
TITLE_FONT = "黑体"
BODY_SIZE = Pt(10.5)
LINE_SPACING = 1.15


def render_docx(
    blocks: Sequence[Block],
    output_path: str,
    paper_size: str = "A4",
    header_text: str = "",
) -> str:
    """把 Block 列表渲染成 .docx，返回写入路径。"""
    doc = Document()
    section = doc.sections[0]
    _setup_page(section, paper_size)
    _setup_default_style(doc)

    if header_text:
        _write_header(section, header_text)

    for block in blocks:
        _write_block(doc, block)

    doc.save(output_path)
    return output_path


# ── 页面 ─────────────────────────────────────────────────────────────


def _setup_page(section, paper_size: str) -> None:
    width, height = PAPER_SIZES.get(paper_size, PAPER_SIZES["A4"])
    if paper_size == "A3":
        width, height = height, width  # A3 横向
        section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(width)
    section.page_height = Mm(height)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)


def _setup_default_style(doc) -> None:
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE
    style.paragraph_format.line_spacing = LINE_SPACING
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(1)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)


def _write_header(section, header_text: str) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = paragraph.add_run(header_text)
    run.font.size = Pt(9)
    _cjk(run, BODY_FONT, Pt(9))


def _cjk(run, font_name: str, size=None) -> None:
    """给 run 设置中西文字体（东亚字体必须单独设 w:eastAsia）。"""
    run.font.name = font_name
    if size is not None:
        run.font.size = size
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def _blank(doc, space_after: Pt = Pt(1), space_before: Pt = Pt(0)):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = LINE_SPACING
    paragraph.paragraph_format.space_before = space_before
    paragraph.paragraph_format.space_after = space_after
    return paragraph


# ── Block 写入 ───────────────────────────────────────────────────────


def _write_block(doc, block: Block) -> None:
    if isinstance(block, PageBreak):
        doc.add_page_break()
        return

    if isinstance(block, Spacer):
        _blank(doc, space_after=Pt(block.height / 2))
        return

    if isinstance(block, Heading):
        _write_heading(doc, block)
        return

    if isinstance(block, SectionHeader):
        paragraph = _blank(doc, space_after=Pt(3), space_before=Pt(6))
        run = paragraph.add_run(block.text)
        run.bold = True
        _cjk(run, BODY_FONT, BODY_SIZE)
        return

    if isinstance(block, Paragraph):
        paragraph = _blank(doc)
        if block.align == "center":
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = paragraph.add_run(block.text)
        run.bold = block.bold
        _cjk(run, BODY_FONT, Pt(block.size))
        return

    if isinstance(block, QuestionBlock):
        _write_question(doc, block)
        return

    if isinstance(block, AnswerKey):
        for group in block.groups:
            paragraph = _blank(doc, space_after=Pt(3))
            run = paragraph.add_run(f"{group.label}答案: ")
            run.bold = True
            _cjk(run, BODY_FONT, BODY_SIZE)
            run = paragraph.add_run(group.span_text)
            _cjk(run, BODY_FONT, BODY_SIZE)

        if block.essay:
            paragraph = _blank(doc, space_after=Pt(3), space_before=Pt(8))
            run = paragraph.add_run("简答题参考答案:")
            run.bold = True
            _cjk(run, BODY_FONT, BODY_SIZE)

            for number, answer in block.essay:
                paragraph = _blank(doc, space_after=Pt(7))
                paragraph.paragraph_format.left_indent = Cm(0.6)
                run = paragraph.add_run(f"{number}. ")
                run.bold = True
                _cjk(run, BODY_FONT, BODY_SIZE)
                run = paragraph.add_run(answer)
                _cjk(run, BODY_FONT, BODY_SIZE)
        return

    if isinstance(block, Instructions):
        if block.title:
            paragraph = _blank(doc, space_after=Pt(3), space_before=Pt(6))
            run = paragraph.add_run(block.title)
            run.bold = True
            _cjk(run, BODY_FONT, BODY_SIZE)
        for line in block.lines:
            paragraph = _blank(doc, space_after=Pt(2))
            run = paragraph.add_run(line)
            _cjk(run, BODY_FONT, Pt(10))
        return

    if isinstance(block, AnswerCardHead):
        _write_answer_card_head(doc, block)
        return

    if isinstance(block, AnswerCardSection):
        _write_answer_card_grid(doc, block)
        return

    if isinstance(block, AnswerCardEssay):
        for number in block.numbers:
            paragraph = _blank(doc, space_after=Pt(0), space_before=Pt(4))
            run = paragraph.add_run(f"{number}.")
            run.bold = True
            _cjk(run, BODY_FONT, BODY_SIZE)
            for _ in range(block.lines_per_question):
                _ruled_paragraph(doc)
            _blank(doc, space_after=Pt(4))
        return


def _write_heading(doc, block: Heading) -> None:
    level = 0 if block.level == 1 else 1
    heading = doc.add_heading(block.text, level=level)
    heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    size = Pt(16) if block.level == 1 else Pt(13)
    for run in heading.runs:
        _cjk(run, TITLE_FONT, size)


def _write_question(doc, block: QuestionBlock) -> None:
    paragraph = _blank(doc)
    run = paragraph.add_run(f"{block.number}.")
    run.bold = True
    _cjk(run, BODY_FONT, BODY_SIZE)

    tail = f" {block.stem}"
    tail += " __________" if block.answer_blank else f" [{block.label}]"
    run = paragraph.add_run(tail)
    _cjk(run, BODY_FONT, BODY_SIZE)

    if block.options:
        paragraph = _blank(doc, space_after=Pt(3))
        paragraph.paragraph_format.left_indent = Cm(0.6)
        rendered = "    ".join(f"{letter}. {text}" for letter, text in block.options)
        run = paragraph.add_run(rendered)
        _cjk(run, BODY_FONT, BODY_SIZE)

    for _ in range(block.blank_lines):
        _ruled_paragraph(doc)


def _ruled_paragraph(doc):
    """一条作答横线。

    用段落下边框而不是一串下划线字符：线连续、无断口，黑白打印稳定。
    """
    paragraph = _blank(doc, space_after=Pt(0))
    paragraph.paragraph_format.line_spacing = 1.6

    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    pPr.append(borders)
    return paragraph


def _write_answer_card_head(doc, block: AnswerCardHead) -> None:
    if block.title:
        _write_heading(doc, Heading(block.title, level=1))
    if block.fields:
        paragraph = _blank(doc, space_after=Pt(6))
        run = paragraph.add_run("　　".join(f"{f}：__________" for f in block.fields))
        _cjk(run, BODY_FONT, Pt(11))
    if block.paper_type_label and block.paper_type_count > 0:
        paragraph = _blank(doc, space_after=Pt(6))
        boxes = "　".join(f"[ ] {i + 1}" for i in range(block.paper_type_count))
        run = paragraph.add_run(f"{block.paper_type_label}：{boxes}")
        _cjk(run, BODY_FONT, Pt(11))


def _write_answer_card_grid(doc, block: AnswerCardSection) -> None:
    """填写式作答区（与 PDF 侧同一套版式）。

    每行 ``ANSWER_CARD_PER_ROW`` 题，两行一组：上面题号、下面留白作答。
    列宽与行高都必须显式设置 —— 交给 Word 自动伸缩的话，表格可能比版心还宽。
    """
    if block.number_end < block.number_start:
        return

    question_bands = bands(block.number_start, block.number_end)
    if not question_bands:
        return

    column_count = max(len(band) for band in question_bands)

    # A4 版心 = 21cm - 左 2cm - 右 2cm = 17cm，再留 2mm 安全余量：
    # python-docx 的宽度要经 EMU/twip 两次换算，紧贴 17.00cm 会因舍入
    # 变成 17.01cm 而溢出到页边距里。
    content_cm = (PAPER_SIZES["A4"][0] - 20.0 * 2 - 2.0) / 10.0
    column_cm = content_cm / column_count

    rows = []
    for band in question_bands:
        number_row = [str(n) for n in band]
        number_row += [""] * (column_count - len(band))
        rows.append(number_row)
        rows.append([""] * column_count)

    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.autofit = False

    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.width = Cm(column_cm)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            if value:
                run = paragraph.add_run(value)
                _cjk(run, BODY_FONT, Pt(9))

    # 列宽要同时写到 tblGrid 与每个单元格，否则部分 Word 版本会按内容重新伸缩
    for index in range(column_count):
        table.columns[index].width = Cm(column_cm)
        for cell in table.columns[index].cells:
            cell.width = Cm(column_cm)

    # 交替行高：题号行矮、作答行高
    for index, row in enumerate(table.rows):
        row.height = Cm(0.65 if index % 2 == 0 else 0.95)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    _blank(doc, space_after=Pt(6))
