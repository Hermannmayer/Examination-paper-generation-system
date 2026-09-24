"""题库读取：xlsx（openpyxl 直读）+ CSV。

取代原先的 ``pandas.read_excel``。**不使用 pandas / numpy**，
因为 openpyxl 自身会 import numpy，打包时必须能把它排除掉。

pandas 的 dtype 推断掩盖了不少问题（整列无缺失 → ``int64`` → ``"1"``；
有缺失 → ``float64`` → ``"1.0"``）。这里的归一化必须显式处理
int / float / str / bool / datetime 五种单元格类型。
"""

from __future__ import annotations

import csv
import os
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

from .constants import OPTION_LETTERS, REQUIRED_COLUMNS
from .errors import Codes, DomainError
from .models import Question, Warning, make_qid

# 可选列的常见别名（首个命中的即为该字段）
OPTIONAL_ALIASES = {
    "difficulty": ["难度", "难度等级", "难易度"],
    "knowledge": ["知识点", "考点", "章节", "知识模块"],
}

EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}
CSV_EXTENSIONS = {".csv", ".txt"}


@dataclass
class LoadResult:
    """一次题库加载的结果。"""

    questions: List[Question] = field(default_factory=list)
    source: str = ""
    sheet: str = ""
    columns: List[str] = field(default_factory=list)
    warnings: List[Warning] = field(default_factory=list)

    @property
    def counts_by_type(self) -> Dict[str, int]:
        out: Dict[str, int] = {}
        for q in self.questions:
            out[q.qtype] = out.get(q.qtype, 0) + 1
        return out

    @property
    def difficulty_values(self) -> List[str]:
        seen: List[str] = []
        for q in self.questions:
            if q.difficulty and q.difficulty not in seen:
                seen.append(q.difficulty)
        return seen

    @property
    def knowledge_values(self) -> List[str]:
        seen: List[str] = []
        for q in self.questions:
            if q.knowledge and q.knowledge not in seen:
                seen.append(q.knowledge)
        return seen

    def to_dict(self) -> Dict[str, Any]:
        return {
            "source": self.source,
            "sheet": self.sheet,
            "columns": list(self.columns),
            "total": len(self.questions),
            "counts_by_type": self.counts_by_type,
            "difficulty_values": self.difficulty_values,
            "knowledge_values": self.knowledge_values,
            "warnings": [w.to_dict() for w in self.warnings],
        }


# ── 单元格标准化 ─────────────────────────────────────────────────────


def cell_to_text(value: Any) -> str:
    """把任意单元格值转成干净字符串。

    关键点：``1.0``(float) 必须变成 ``"1"`` 而不是 ``"1.0"``，
    否则判断题答案会渲染成 "1.0"。
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _clean_option_text(text: Any) -> str:
    """清理选项文本：去除 ``[A]`` 前缀（由 ``core.py:_clean_option_text`` 迁移）。"""
    t = cell_to_text(text)
    if len(t) > 2 and t.startswith("[") and t[1:2].isalpha() and t[2:3] == "]":
        return t[3:].strip()
    return t


_TRUTHY = {"1", "√", "对", "是", "正确", "T", "TRUE", "Y", "YES", "V"}
_FALSY = {"0", "×", "错", "否", "错误", "F", "FALSE", "N", "NO", "X"}


def normalize_answer(value: Any, qtype: str) -> str:
    """答案归一化。

    判断题的 ``1`` / ``1.0`` / ``"1"`` / ``"√"`` / ``"对"`` 统一成 ``√``；
    ``0`` 系列统一成 ``×``。其他题型原样返回。
    """
    text = cell_to_text(value)
    if not text:
        return ""

    if qtype == "判断题":
        upper = text.upper()
        if text in _TRUTHY or upper in _TRUTHY:
            return "√"
        if text in _FALSY or upper in _FALSY:
            return "×"
    return text


def _is_blank(value: Any) -> bool:
    """替代 ``pandas.notna`` —— 空单元格判定。"""
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return False


# ── 表头解析 ─────────────────────────────────────────────────────────


def _normalize_header(value: Any) -> str:
    """表头去 BOM、去空白、去零宽字符。"""
    text = cell_to_text(value)
    return text.replace("﻿", "").replace("​", "").strip()


def _resolve_columns(header_row: List[Any]) -> Dict[str, int]:
    """返回 {字段名: 列下标}。字段名含必填列与可选别名命中的列。"""
    normalized = [_normalize_header(v) for v in header_row]
    mapping: Dict[str, int] = {}

    for name in REQUIRED_COLUMNS:
        if name in normalized:
            mapping[name] = normalized.index(name)
    for letter in OPTION_LETTERS:
        col = f"选项{letter}"
        if col in normalized:
            mapping[col] = normalized.index(col)

    for field_name, aliases in OPTIONAL_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                mapping[field_name] = normalized.index(alias)
                break

    return mapping


def _detect_sheet(workbook, requested: Optional[str]):
    if requested:
        if requested not in workbook.sheetnames:
            raise DomainError(
                Codes.EMPTY_SHEET,
                f"工作表不存在：{requested}",
                hint=f"可用的工作表：{'、'.join(workbook.sheetnames)}",
                details={"available": list(workbook.sheetnames)},
            )
        return workbook[requested]
    return workbook[workbook.sheetnames[0]]


def _build_questions(rows, mapping, source, sheet) -> List[Question]:
    """把行数据转成 Question 列表。``rows`` 是已跳表头的行迭代器。"""
    stem_col = mapping["题目"]
    answer_col = mapping["正确答案"]
    type_col = mapping["题型"]

    questions: List[Question] = []
    seen_qids: Dict[str, int] = {}

    for row_number, row in rows:
        values = list(row)
        if all(_is_blank(v) for v in values):
            continue  # 整行为空

        qtype = cell_to_text(values[type_col]) if type_col < len(values) else ""
        stem = cell_to_text(values[stem_col]) if stem_col < len(values) else ""

        if not qtype and not stem:
            continue
        if not stem:
            raise DomainError(
                Codes.EMPTY_SHEET,
                f"第 {row_number} 行缺少题目内容",
                hint="请检查题库中是否有空行或错位的数据",
                details={"row": row_number},
            )

        raw_answer = values[answer_col] if answer_col < len(values) else None
        answer = normalize_answer(raw_answer, qtype)

        options = []
        for letter in OPTION_LETTERS:
            key = f"选项{letter}"
            idx = mapping.get(key)
            if idx is None or idx >= len(values):
                continue
            text = _clean_option_text(values[idx])
            if text:
                options.append((letter, text))

        difficulty = None
        if "difficulty" in mapping and mapping["difficulty"] < len(values):
            difficulty = cell_to_text(values[mapping["difficulty"]]) or None

        knowledge = None
        if "knowledge" in mapping and mapping["knowledge"] < len(values):
            knowledge = cell_to_text(values[mapping["knowledge"]]) or None

        qid = make_qid(qtype, stem, answer, options)
        if qid in seen_qids:
            # 同一题库内的完全重复题：保留第一条，后续丢弃
            seen_qids[qid] += 1
            continue
        seen_qids[qid] = 1

        questions.append(
            Question(
                qid=qid,
                qtype=qtype,
                stem=stem,
                answer_raw=cell_to_text(raw_answer),
                answer=answer,
                options=tuple(options),
                difficulty=difficulty,
                knowledge=knowledge,
                source_row=row_number,
                source_sheet=sheet,
            )
        )

    return questions


def _finalize(questions, source, sheet, header_row, warnings) -> LoadResult:
    if not questions:
        raise DomainError(
            Codes.EMPTY_SHEET,
            "题库中没有读到任何题目",
            hint="请确认第一行是表头、从第二行开始是题目数据",
            details={"source": source, "sheet": sheet},
        )
    return LoadResult(
        questions=questions,
        source=source,
        sheet=sheet,
        columns=[_normalize_header(v) for v in header_row],
        warnings=warnings,
    )


def _validate_columns(header_row, source, sheet):
    """校验必需列。**在读数据之前**执行，避免读了一半才发现缺列。"""
    mapping = _resolve_columns(header_row)
    missing = [c for c in REQUIRED_COLUMNS if c not in mapping]
    if missing:
        raise DomainError(
            Codes.MISSING_COLUMNS,
            f"题库缺少必需的列：{'、'.join(missing)}",
            hint="请对照「题库模板」补齐列名：「题型」「题目」「正确答案」为必填",
            details={
                "missing": missing,
                "found": [_normalize_header(v) for v in header_row],
                "source": source,
                "sheet": sheet,
            },
        )
    return mapping


# ── 入口 ─────────────────────────────────────────────────────────────


def load_question_bank(path: str, sheet: Optional[str] = None) -> LoadResult:
    """读取题库文件。失败时抛 :class:`DomainError`，**不产生任何部分状态**。"""
    if not path:
        raise DomainError(Codes.FILE_NOT_FOUND, "请先选择题库文件")
    if not os.path.exists(path):
        raise DomainError(
            Codes.FILE_NOT_FOUND,
            f"文件不存在：{path}",
            details={"path": path},
        )

    ext = os.path.splitext(path)[1].lower()
    if ext in EXCEL_EXTENSIONS:
        return _load_excel(path, sheet)
    if ext in CSV_EXTENSIONS:
        return _load_csv(path)
    if ext == ".xls":
        raise DomainError(
            Codes.UNSUPPORTED_FORMAT,
            "不支持 .xls（Excel 97-2003）格式",
            hint="请用 Excel 另存为 .xlsx 后重试",
            details={"path": path, "extension": ext},
        )
    raise DomainError(
        Codes.UNSUPPORTED_FORMAT,
        f"不支持的文件格式：{ext or '（无扩展名）'}",
        hint="支持 .xlsx / .xlsm / .csv",
        details={"path": path, "extension": ext},
    )


def _load_excel(path: str, sheet: Optional[str]) -> LoadResult:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # 打开失败：损坏文件、被占用、加密等
        raise DomainError(
            Codes.UNSUPPORTED_FORMAT,
            f"无法打开题库文件：{exc}",
            hint="请确认文件未损坏、未被其他程序占用、且未加密",
            details={"path": path},
        ) from exc

    try:
        worksheet = _detect_sheet(workbook, sheet)
        rows = worksheet.iter_rows(values_only=True)

        try:
            header_row = list(next(rows))
        except StopIteration:
            raise DomainError(
                Codes.EMPTY_SHEET,
                "题库工作表是空的",
                details={"path": path, "sheet": worksheet.title},
            ) from None

        mapping = _validate_columns(header_row, path, worksheet.title)

        numbered = ((idx, row) for idx, row in enumerate(rows, start=2))
        questions = _build_questions(numbered, mapping, path, worksheet.title)
        return _finalize(questions, path, worksheet.title, header_row, [])
    finally:
        workbook.close()


def _load_csv(path: str) -> LoadResult:
    text, encoding = _read_text_with_fallback(path)
    delimiter = _sniff_delimiter(text)

    reader = csv.reader(text.splitlines(), delimiter=delimiter)
    try:
        header_row = next(reader)
    except StopIteration:
        raise DomainError(
            Codes.EMPTY_SHEET, "CSV 文件是空的", details={"path": path}
        ) from None

    mapping = _validate_columns(header_row, path, "(csv)")

    warnings: List[Warning] = []
    if encoding.lower() not in ("utf-8", "utf-8-sig"):
        warnings.append(
            Warning(
                "ENCODING_FALLBACK",
                f"CSV 不是 UTF-8 编码，已按 {encoding} 读取",
                {"encoding": encoding},
            )
        )

    numbered = ((idx, row) for idx, row in enumerate(reader, start=2))
    questions = _build_questions(numbered, mapping, path, "(csv)")
    return _finalize(questions, path, "(csv)", header_row, warnings)


def _read_text_with_fallback(path: str):
    for encoding in ("utf-8-sig", "utf-8", "gbk", "gb18030", "big5"):
        try:
            with open(path, "r", encoding=encoding, newline="") as fh:
                return fh.read(), encoding
        except UnicodeDecodeError:
            continue
    raise DomainError(
        Codes.UNSUPPORTED_FORMAT,
        "无法识别 CSV 文件的编码",
        hint="请将文件另存为 UTF-8 或 GBK 编码",
        details={"path": path},
    )


def _sniff_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:5])
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","
