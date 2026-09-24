"""Word 渲染测试。

与 PDF 消费同一套 Block，因此题目顺序、题号、答案分组必然一致。
"""

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT

from app.docmodel import build_paper_blocks
from conftest import build_paper, count_page_breaks


def _render(result, cfg, tmp_path, paper_size="A4", name="out.docx"):
    from renderers.docx_renderer import render_docx

    blocks = build_paper_blocks(result.papers[0], cfg)
    path = tmp_path / name
    render_docx(blocks, str(path), paper_size, cfg.header_text)
    return path, Document(str(path))


def _all_text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def test_questions_and_numbers_are_present(paper_bundle, tmp_path):
    result, _pool, cfg = paper_bundle
    _, doc = _render(result, cfg, tmp_path)

    text = _all_text(doc)
    assert "判断题" in text
    assert "单选题" in text
    # 题号是加粗 run，但文本仍在
    assert "1." in text


def test_no_page_break_when_answers_excluded(paper_bundle, tmp_path):
    """不含答案时不应有分页符（否则末尾多一张空白页）。"""
    result, _pool, cfg = paper_bundle
    cfg.include_answers = False

    path, _ = _render(result, cfg, tmp_path)
    assert count_page_breaks(path) == 0


def test_page_break_before_answers(paper_bundle, tmp_path):
    result, _pool, cfg = paper_bundle
    cfg.include_answers = True

    path, doc = _render(result, cfg, tmp_path)
    assert count_page_breaks(path) == 1
    assert "参考答案" in _all_text(doc)


def test_answer_groups_use_real_numbers(paper_bundle, tmp_path):
    result, _pool, cfg = paper_bundle
    cfg.include_answers = True

    _, doc = _render(result, cfg, tmp_path)
    text = _all_text(doc)
    assert "全部答案" in text
    assert "判断题答案" in text


def test_a3_is_landscape(paper_bundle, tmp_path):
    result, _pool, cfg = paper_bundle
    _, doc = _render(result, cfg, tmp_path, paper_size="A3")

    section = doc.sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    # python-docx 以 twip 存储页面尺寸，读回来有亚毫米级舍入，
    # 所以按毫米取整比较而不是逐 EMU 相等。
    assert round(section.page_width.mm) == 420
    assert round(section.page_height.mm) == 297


def test_header_text_is_written(bank_path, tmp_path):
    result, _pool, cfg = build_paper(bank_path, header_text="某校期末考试")
    _, doc = _render(result, cfg, tmp_path)

    header_text = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
    assert "某校期末考试" in header_text


def test_stem_starting_with_digit_is_not_misparsed(tmp_path):
    """题干以数字开头时，题号仍应正确 —— 旧实现在 docx 层用
    ``line[0].isdigit()`` 反解析纯文本，会把这种题判错。"""
    import openpyxl

    bank = tmp_path / "digit.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["题型", "题目", "正确答案"])
    ws.append(["判断题", "2026年养老护理新规要求……", "1"])
    wb.save(bank)

    result, _pool, cfg = build_paper(bank)
    _, doc = _render(result, cfg, tmp_path)

    text = _all_text(doc)
    assert "2026年养老护理新规要求" in text
    assert "1. 2026年" in text.replace(" ", " ")
