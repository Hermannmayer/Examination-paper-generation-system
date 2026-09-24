"""答题卡测试。

采用填写式排版（写 √/× 或 ABCD），核心诉求是**省地方**：
满题库 120 题也要压进一页 A4，同时题号必须与试卷一致。
"""

from pathlib import Path

import pytest
from pypdf import PdfReader

from app.docmodel import AnswerCardHead, AnswerCardSection, Instructions
from renderers.answer_card import build_answer_card_blocks, build_answer_card_document
from renderers.pdf_renderer import render_pdf

# 仓库自带的示例题库（120 题），用于满量排版测试
BANK = Path(__file__).resolve().parent.parent / "题库模板.xlsx"


def _blocks(paper_bundle):
    result, _pool, cfg = paper_bundle
    return build_answer_card_blocks(result.papers[0], cfg, paper_count=1)


def test_card_structure(paper_bundle):
    blocks = _blocks(paper_bundle)
    kinds = [type(b).__name__ for b in blocks]

    assert kinds[0] == "AnswerCardHead"
    assert "AnswerCardSection" in kinds
    assert "Instructions" in kinds


def test_card_sections_match_paper_question_numbers(paper_bundle):
    """答题卡的题号必须与试卷一致。"""
    result, _pool, cfg = paper_bundle
    paper = result.papers[0]
    blocks = _blocks(paper_bundle)

    grids = [b for b in blocks if isinstance(b, AnswerCardSection)]
    assert len(grids) == len([s for s in paper.sections if s.count > 0])

    for grid, section in zip(grids, [s for s in paper.sections if s.count > 0]):
        assert grid.number_start == section.number_start
        assert grid.number_end == section.number_end
        assert grid.label == section.label


def test_judgment_hint_is_check_and_cross(paper_bundle):
    blocks = _blocks(paper_bundle)
    judgment = [b for b in blocks if isinstance(b, AnswerCardSection) and b.label == "判断题"]
    assert judgment, "题库里有判断题"
    assert judgment[0].hint == "√ / ×"


def test_choice_hint_lists_letters(paper_bundle):
    blocks = _blocks(paper_bundle)
    mcq = [b for b in blocks if isinstance(b, AnswerCardSection) and b.label == "单选题"]
    assert mcq
    assert "A" in mcq[0].hint


def test_bands_split_evenly():
    from app.docmodel import ANSWER_CARD_PER_ROW
    from renderers.answer_card import bands

    result = bands(1, 40)
    assert result[0] == list(range(1, ANSWER_CARD_PER_ROW + 1))
    assert result[-1][-1] == 40
    assert all(len(band) <= ANSWER_CARD_PER_ROW for band in result)
    assert [n for band in result for n in band] == list(range(1, 41))


def test_bands_handles_single_question():
    from renderers.answer_card import bands

    assert bands(7, 7) == [[7]]
    assert bands(5, 4) == []


def test_120_questions_fit_on_one_page(tmp_path):
    """填写式排版的核心目的：满题库也要压进一页 A4。

    旧的涂卡式排 120 题要两三页，是用户明确要求改掉的问题。
    """
    import pymupdf

    from app.generator import generate
    from app.loader import load_question_bank
    from app.models import PaperConfig
    from app.pool import QuestionPool
    from renderers.answer_card import build_answer_card_document
    from renderers.pdf_renderer import render_pdf

    pool = QuestionPool(load_question_bank(str(BANK)))
    config = PaperConfig.from_legacy({
        "export_mode": "随机抽取",
        "exam_title": "满题库测试",
        "include_judgment": True, "include_mcq": True, "include_mcq_multi": True,
        "judgment_count": 40, "mcq_count": 50, "mcq_multi_count": 30,
    })
    result = generate(config, pool)
    assert result.papers[0].total_count == 120

    path = tmp_path / "card.pdf"
    render_pdf(
        build_answer_card_document(result.papers, config), str(path), "A4", ""
    )

    reader = pymupdf.open(str(path))
    assert reader.page_count == 1, f"120 题应压进 1 页，实际 {reader.page_count} 页"

    page = reader[0]
    spans = [
        span["bbox"]
        for block in page.get_text("dict")["blocks"]
        for line in block.get("lines", [])
        for span in line["spans"]
    ]
    overflow = [
        box for box in spans
        if box[0] < 56 or box[2] > 539 or box[3] > page.rect.height - 42
    ]
    assert not overflow, f"有 {len(overflow)} 处内容超出可打印区域"


def test_card_header_has_identity_fields(paper_bundle):
    blocks = _blocks(paper_bundle)
    head = blocks[0]
    assert isinstance(head, AnswerCardHead)
    assert "姓名" in head.fields
    assert "考号" in head.fields


def test_multi_paper_card_marks_paper_type(paper_bundle):
    result, _pool, cfg = paper_bundle
    blocks = build_answer_card_blocks(result.papers[0], cfg, paper_count=4)
    head = blocks[0]
    assert head.paper_type_count == 4


def test_card_renders_to_pdf(paper_bundle, tmp_path):
    blocks = _blocks(paper_bundle)
    path = tmp_path / "card.pdf"
    render_pdf(blocks, str(path), "A4", "")

    reader = PdfReader(str(path))
    text = "".join(p.extract_text() for p in reader.pages)
    assert "答题卡" in text
    assert "注意事项" in text
    assert "姓名" in text


def test_multi_paper_cards_start_new_pages(bank_path, tmp_path):
    from conftest import build_paper

    result, _pool, cfg = build_paper(bank_path, exam_count=3)
    blocks = build_answer_card_document(result.papers, cfg)
    path = tmp_path / "cards.pdf"
    render_pdf(blocks, str(path), "A4", "")

    reader = PdfReader(str(path))
    assert len(reader.pages) >= 3
    for page in reader.pages:
        assert "答题卡" in page.extract_text()


def test_card_renders_to_docx(paper_bundle, tmp_path):
    from docx import Document

    from renderers.docx_renderer import render_docx

    path = tmp_path / "card.docx"
    render_docx(_blocks(paper_bundle), str(path), "A4", "")

    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "答题卡" in text
    assert doc.tables, "涂卡区应渲染成表格"


# ── 出答题卡时，试卷上不再留作答空间 ─────────────────────────────────


def _paper_for(_tmp_path, include_card):
    """用仓库自带题库（含判断题与简答题）组一份卷，返回试卷的 Block。"""
    from app.docmodel import build_paper_blocks
    from app.generator import generate
    from app.loader import load_question_bank
    from app.models import PaperConfig
    from app.pool import QuestionPool

    pool = QuestionPool(load_question_bank(str(BANK)))
    config = PaperConfig.from_legacy({
        "export_mode": "随机抽取",
        "exam_title": "作答空间测试",
        "include_judgment": True,
        "include_short_answer": True,
        "judgment_count": 3,
        "short_answer_count": 2,
        "include_answer_card": include_card,
    })
    paper = generate(config, pool).papers[0]
    return build_paper_blocks(paper, config), config, paper


def test_paper_keeps_write_space_without_answer_card(tmp_path):
    """不出答题卡时，学生只能写在试卷上，横线必须留。"""
    from app.docmodel import QuestionBlock

    blocks, _config, _paper = _paper_for(tmp_path, include_card=False)

    judgment = [b for b in blocks if isinstance(b, QuestionBlock) and b.label == "判断题"]
    essay = [b for b in blocks if isinstance(b, QuestionBlock) and b.label == "简答题"]

    assert judgment and all(b.answer_blank for b in judgment), "判断题应留作答线"
    assert essay and all(b.blank_lines > 0 for b in essay), "简答题应留作答横线"


def test_answer_card_removes_write_space_from_paper(tmp_path):
    """出了答题卡，学生写在卡上，试卷上再留横线就是白占版面。"""
    from app.docmodel import QuestionBlock

    blocks, _config, _paper = _paper_for(tmp_path, include_card=True)

    judgment = [b for b in blocks if isinstance(b, QuestionBlock) and b.label == "判断题"]
    essay = [b for b in blocks if isinstance(b, QuestionBlock) and b.label == "简答题"]

    assert judgment and not any(b.answer_blank for b in judgment), "判断题不该再留作答线"
    assert essay and all(b.blank_lines == 0 for b in essay), "简答题不该再留作答横线"


def test_answer_card_still_provides_write_area(tmp_path):
    """试卷上不留了，但答题卡上必须留 —— 否则学生无处作答。"""
    from app.docmodel import AnswerCardEssay, QuestionBlock

    _blocks_, config, paper = _paper_for(tmp_path, include_card=True)
    card_blocks = build_answer_card_blocks(paper, config, paper_count=1)

    essay_areas = [b for b in card_blocks if isinstance(b, AnswerCardEssay)]
    assert essay_areas, "答题卡上应有简答题作答区"
    assert all(area.lines_per_question > 0 for area in essay_areas)
    assert all(len(area.numbers) > 0 for area in essay_areas)


def _section_lines(content, marker):
    """取某个小节标题之后、到第一个空行为止的题目行。

    注意不能直接搜 ``__________`` —— 考生信息那一行
    （姓名：__________  考号：__________）里也有，会误判。
    """
    start = content.index(marker)
    lines = []
    started = False
    for line in content[start:].split("\n")[1:]:
        text = line.strip()
        if not text:
            if started:      # 已经收过题了，遇到空行说明本节结束
                break
            continue         # 标题与题目之间本来就有空行
        started = True
        lines.append(text)
    return lines


def test_text_render_follows_the_same_rule(tmp_path):
    """纯文本渲染（paper.content）必须和 Block 渲染遵守同一规则。"""
    _blocks_a, _config_a, paper_with_card = _paper_for(tmp_path, include_card=True)
    _blocks_b, _config_b, paper_without_card = _paper_for(tmp_path, include_card=False)

    with_card = _section_lines(paper_with_card.content, "判断题（")
    without_card = _section_lines(paper_without_card.content, "判断题（")

    assert with_card and without_card, "两份卷子都应有判断题小节"
    assert all(line.endswith("__________") for line in without_card), (
        f"没有答题卡时判断题应留作答线，实际：{without_card}"
    )
    assert not any(line.endswith("__________") for line in with_card), (
        f"有答题卡时判断题不该留作答线，实际：{with_card}"
    )
